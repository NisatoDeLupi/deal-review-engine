"""
run_enrichment.py — reference runner that drives a full enrichment.

This file is illustrative. Most real runs are written ad-hoc per workbook
because per-row company summaries (Phase 2 input) come from Claude, not from
disk. But the shape below is the recommended pattern:

  1. Load the workbook once.
  2. For each sheet that needs enrichment:
       a. Read every row into memory as a dict (header -> value).
       b. Phase 2 batches of 25 — for each batch:
            - Have Claude produce {row_num: {summary, divisions}} for the batch
              (this is the only LLM-touching part).
            - Call enrich_engine.build_enrichment for each row (Phase 1 logic
              runs here, alongside the AI summary going through unchanged).
            - Call enrich_engine.write_batch(...) to save the batch to disk
              before moving to the next batch.

Usage sketch (pseudo-code — the LLM call is a placeholder):

    import openpyxl
    from enrich_engine import build_enrichment, write_batch

    OUT = "Enriched.xlsx"
    wb = openpyxl.load_workbook(OUT)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        headers = [c.value for c in ws[1]]
        rows = [
            (r, {headers[c-1]: ws.cell(row=r, column=c).value
                 for c in range(1, len(headers)+1)})
            for r in range(2, ws.max_row + 1)
        ]
        for batch_start in range(0, len(rows), 25):
            batch = rows[batch_start:batch_start + 25]
            ai = ask_claude_for_summaries(batch)  # returns {row_num: {summary, divisions}}
            enriched = []
            for row_num, row in batch:
                meta = ai.get(row_num, {})
                enr = build_enrichment(
                    row,
                    row.get("Company Name") or f"Row {row_num}",
                    meta.get("summary", "Insufficient public information."),
                    meta.get("divisions", "No major divisions identified."),
                )
                enriched.append((row_num, enr))
            write_batch(OUT, sheet_name, enriched)
            print(f"[{sheet_name}] saved rows {batch[0][0]}..{batch[-1][0]}")

Keep the LLM call narrow — it owns Company Summary and Sister Verticals only.
Everything else is build_enrichment's job, which keeps the run reproducible.
"""
