#!/usr/bin/env python3
"""Build the deal-review workbook + HubSpot note/task posting payload from master.json."""
import json, datetime, html, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

MASTER = "/tmp/dr/master.json"
OUT_XLSX = "/tmp/dr/Deal_Review_All_99_Deals.xlsx"
OUT_PAYLOAD = "/tmp/dr/posting_payload.json"

MANAGER_ID = "900000011"; MANAGER_NAME = "Abhimanyu Shekhawat"
AE_NAMES = {"900000113":"AE_13","900000115":"AE_27","900000080":"SDR_Q","900000015":"AE_74"}
FX = {"INR":1,"USD":90,"AED":25,"SGD":68,"EUR":96,"GBP":113}
STAGE = {  # value -> (label, short, rank)
 "appointmentscheduled":("1 - Prospect","Prospect",1),"90000040":("1 - Prospect","Prospect",1),
 "900000059":("1 - Prospect","Prospect",1),"900000050":("1 - Prospect","Prospect",1),
 "9000000003":("2 - Need Analysis & Solutioning","Need Analysis",2),"9000000004":("2 - Need Analysis & Solutioning","Need Analysis",2),
 "900000051":("2 - Qualified","Qualified",2),
 "presentationscheduled":("3 - Opportunity","Opportunity",3),"90000042":("3 - Opportunity","Opportunity",3),
 "900000060":("2 - Opportunity","Opportunity",3),"900000052":("3 - Opportunity","Opportunity",3),
 "decisionmakerboughtin":("4 - Proposal","Proposal",4),"90000043":("4 - Proposal","Proposal",4),
 "900000061":("3 - Proposal","Proposal",4),"900000053":("4 - Proposal","Proposal",4),
 "contractsent":("5 - Negotiation","Negotiation",5),"90000044":("5 - Negotiation","Negotiation",5),
 "900000062":("4 - Negotiation","Negotiation",5),"900000054":("5 - Negotiation","Negotiation",5),
}
STAGE_FILL={"Negotiation":"C6EFCE","Proposal":"FFEB9C","Opportunity":"FFC7CE","Need Analysis":"DDEBF7","Qualified":"DDEBF7","Prospect":"D9D9D9"}
AE_FILL={"Yes":"C6EFCE","Partial":"FFEB9C","No":"FFC7CE"}
ACT_FILL={"Active":"C6EFCE","One-way":"FFEB9C","Silent":"FFC7CE","Dead":"C00000"}
TODAY=datetime.date(2026,6,4)

def add_bd(d,n):
    while n>0:
        d+=datetime.timedelta(days=1)
        if d.weekday()<5: n-=1
    return d
AE_DUE=add_bd(TODAY,1)          # Fri Jun 5
MGR_DUE=add_bd(AE_DUE,4)        # Thu Jun 11
AE_DUE_ISO=AE_DUE.strftime("%Y-%m-%dT11:30:00Z")
MGR_DUE_ISO=MGR_DUE.strftime("%Y-%m-%dT11:30:00Z")

def parse_dt(s):
    if not s: return None
    try: return datetime.datetime.fromisoformat(s.replace("Z","+00:00")).date()
    except: return None

def product(name):
    n=(name or "").lower()
    for p in ["empuls","compass","plum","loyalife"]:
        if p in n: return p.capitalize()
    if "api" in n or "rewards" in n: return "Plum"
    if "survey" in n: return "Empuls"
    return "—"

def inr(amount,cur):
    try: a=float(amount or 0)
    except: a=0
    return a*FX.get(cur,1)

def flatten(lst):
    out=[]
    for it in (lst or []):
        if isinstance(it,dict):
            t=it.get("type") or it.get("category") or ""
            txt=it.get("text") or it.get("detail") or it.get("quote") or it.get("signal") or it.get("objection") or ""
            out.append((f"{t}: {txt}".strip(": ").strip()) or json.dumps(it))
        else: out.append(str(it))
    return out

deals=json.load(open(MASTER))
for d in deals:
    d["_stage"]=STAGE.get(d.get("dealstage"),("?",str(d.get("dealstage")),0))
    d["_inr"]=inr(d.get("amount"),d.get("cur"))
    d["_prod"]=product(d.get("dealname"))
    dates=[parse_dt(d.get(k)) for k in ("last_email_date","last_call_date","last_note_date")]
    dates=[x for x in dates if x]
    d["_last"]=max(dates) if dates else None
    d["_idle"]=(TODAY-d["_last"]).days if d["_last"] else None

# ---------- posting payload ----------
def esc(t): return html.escape(t or "").replace("\n","<br>")
def note_html(d):
    aeid=d["owner"]; aename=AE_NAMES.get(aeid,aeid)
    summ=d.get("ae_summary") or ""
    if d.get("ae_summary_status")=="No" or not summ.strip():
        summ="AE SUMMARY MISSING - rep must add."
    mgr=(d.get("abhimanyu_note") or "").strip()
    nxt=(d.get("next_action") or "").strip()
    body=[]
    if mgr: body.append(esc(mgr))
    if nxt: body.append("Next action: "+esc(nxt))
    notes_block="</p><p>".join(body) if body else "Reviewed."
    return (f'<div><p><span data-mention-id="{aeid}" data-mention-name="{aename}" '
        f'style="color: #425b76;font-weight: 600;">@{aename}</span></p><p><br></p>'
        f'<p><strong>DEAL SUMMARY</strong></p><p>{esc(summ)}</p><p><br></p>'
        f'<p><strong>ABHIMANYU NOTES</strong></p><p>{notes_block}</p><p><br></p>'
        f'<p>cc: <span data-mention-id="{MANAGER_ID}" data-mention-name="{MANAGER_NAME}" '
        f'style="color: #425b76;font-weight: 600;">@{MANAGER_NAME}</span></p></div>')

payload=[]
for d in deals:
    did=d["deal_id"]; aeid=d["owner"]; aename=AE_NAMES.get(aeid,aeid)
    pr=(d.get("priority") or "MEDIUM").upper()
    action=(d.get("next_action") or "Review and update deal status.").strip()
    payload.append({
        "deal_id":did,"dealname":d.get("dealname"),"ae":aename,"priority":pr,
        "note":{"objectType":"notes","associations":[{"targetObjectId":int(did),"targetObjectType":"deals"}],
            "properties":{"hs_note_body":note_html(d),"hs_timestamp":TODAY.strftime("%Y-%m-%dT09:00:00Z"),"hubspot_owner_id":MANAGER_ID}},
        "ae_task":{"objectType":"tasks","associations":[{"targetObjectId":int(did),"targetObjectType":"deals"}],
            "properties":{"hs_task_subject":f"Action required: {d.get('dealname')}","hs_task_body":action,
                "hs_task_priority":pr,"hs_task_status":"NOT_STARTED","hs_task_type":"TODO",
                "hs_timestamp":AE_DUE_ISO,"hubspot_owner_id":aeid}},
        "mgr_task":{"objectType":"tasks","associations":[{"targetObjectId":int(did),"targetObjectType":"deals"}],
            "properties":{"hs_task_subject":f"Review: {d.get('dealname')} — verify AE action",
                "hs_task_body":f"AE ({aename}) was tasked to: {action} Confirm completed. If not, escalate or reassign.",
                "hs_task_priority":pr,"hs_task_status":"NOT_STARTED","hs_task_type":"TODO",
                "hs_timestamp":MGR_DUE_ISO,"hubspot_owner_id":MANAGER_ID}},
    })
json.dump(payload,open(OUT_PAYLOAD,"w"),indent=1)

# ---------- workbook ----------
FONT="Arial"
wb=Workbook()
thin=Side(style="thin",color="BFBFBF"); border=Border(thin,thin,thin,thin)
hdr_fill=PatternFill("solid",fgColor="305496"); hdr_font=Font(name=FONT,bold=True,color="FFFFFF",size=10)
def style_hdr(ws,row=1):
    for c in ws[row]:
        c.fill=hdr_fill; c.font=hdr_font; c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); c.border=border
def setw(ws,widths):
    for col,w in widths.items(): ws.column_dimensions[col].width=w

# Sheet 1: Summary
ws=wb.active; ws.title="Summary"
ws["A1"]="Hunting Team Pipeline Review — Open Deals"; ws["A1"].font=Font(name=FONT,bold=True,size=14)
ws["A2"]=f"Run date 2026-06-04 · 99 open deals (stages 1-Prospect → 5-Negotiation) · AE task due {AE_DUE} · Manager review {MGR_DUE}"
ws["A2"].font=Font(name=FONT,italic=True,size=9,color="666666")
r=4; ws.cell(r,1,"AE Scorecard").font=Font(name=FONT,bold=True,size=11); r+=1
heads=["AE","Deals","Pipeline (₹)","AE Summary Yes","Partial","No","Active","One-way","Silent","Dead","HIGH pri"]
for i,h in enumerate(heads,1): ws.cell(r,i,h)
style_hdr(ws,r); hdr_row=r; r+=1
aes=["AE_13","AE_27","SDR_Q","AE_74"]
def cnt(ds,k,v): return sum(1 for x in ds if x.get(k)==v)
for ae in aes:
    ds=[x for x in deals if AE_NAMES.get(x["owner"])==ae]
    vals=[ae,len(ds),round(sum(x["_inr"] for x in ds)),
        cnt(ds,"ae_summary_status","Yes"),cnt(ds,"ae_summary_status","Partial"),cnt(ds,"ae_summary_status","No"),
        cnt(ds,"activity_status","Active"),cnt(ds,"activity_status","One-way"),cnt(ds,"activity_status","Silent"),cnt(ds,"activity_status","Dead"),
        sum(1 for x in ds if (x.get("priority") or "").upper()=="HIGH")]
    for i,v in enumerate(vals,1):
        c=ws.cell(r,i,v); c.border=border; c.font=Font(name=FONT,size=10)
        if i==3: c.number_format='#,##0'
    r+=1
# totals
tr=r; ws.cell(r,1,"TOTAL").font=Font(name=FONT,bold=True)
for i in range(2,12):
    col=chr(64+i)
    ws.cell(r,i,f"=SUM({col}{hdr_row+1}:{col}{r-1})").font=Font(name=FONT,bold=True)
    ws.cell(r,i).border=border
    if i==3: ws.cell(r,i).number_format='#,##0'
ws.cell(r,1).border=border
setw(ws,{"A":20,"B":7,"C":15,"D":13,"E":8,"F":6,"G":8,"H":9,"I":8,"J":7,"K":9})

# critical attention list
r+=2; ws.cell(r,1,"Critical attention (HIGH priority + No-summary or Dead)").font=Font(name=FONT,bold=True,size=11); r+=1
for h in ["Deal","AE","Stage","₹","AE Summary","Activity","Why / Next action"]:
    ws.cell(r,heads.index if False else ["Deal","AE","Stage","₹","AE Summary","Activity","Why / Next action"].index(h)+1,h)
style_hdr(ws,r); r+=1
crit=[x for x in deals if (x.get("priority") or "").upper()=="HIGH" or x.get("ae_summary_status")=="No" or x.get("activity_status")=="Dead"]
crit.sort(key=lambda x:(-x["_inr"]))
for x in crit:
    row=[x.get("dealname"),AE_NAMES.get(x["owner"]),x["_stage"][1],round(x["_inr"]),x.get("ae_summary_status"),x.get("activity_status"),(x.get("next_action") or "")[:160]]
    for i,v in enumerate(row,1):
        c=ws.cell(r,i,v); c.border=border; c.font=Font(name=FONT,size=9); c.alignment=Alignment(wrap_text=True,vertical="top")
        if i==4: c.number_format='#,##0'
        if i==5 and v in AE_FILL: c.fill=PatternFill("solid",fgColor=AE_FILL[v])
        if i==6 and v in ACT_FILL:
            c.fill=PatternFill("solid",fgColor=ACT_FILL[v])
            if v=="Dead": c.font=Font(name=FONT,size=9,color="FFFFFF")
    r+=1

# Sheet 2: Deals
ws2=wb.create_sheet("Deals")
cols=["Deal ID","Deal Name","AE","Stage","Product","Cur","Amount","Amount (₹)","Created","Last Activity","Days Idle",
      "Emails (in/out)","Calls (conn/att)","Last Email","Last Call","AE Summary","AE Status","Activity","Buying Signals","Objections","Auth?","Next Action","Priority","HubSpot"]
ws2.append(cols); style_hdr(ws2)
for x in sorted(deals,key=lambda d:(AE_NAMES.get(d["owner"]),-d["_stage"][2],-d["_inr"])):
    link=f'https://app.hubspot.com/contacts/1234567/record/0-3/{x["deal_id"]}'
    row=[x["deal_id"],x.get("dealname"),AE_NAMES.get(x["owner"]),x["_stage"][1],x["_prod"],x.get("cur"),
        float(x.get("amount") or 0),round(x["_inr"]),(x.get("createdate") or "")[:10],
        x["_last"].isoformat() if x["_last"] else "",x["_idle"],
        f'{x.get("emails_in",0)}/{x.get("emails_out",0)}',f'{x.get("calls_connected",0)}/{x.get("calls_attempted",0)}',
        (x.get("last_email_date") or "")[:10],(x.get("last_call_date") or "")[:10],
        x.get("ae_summary") or "",x.get("ae_summary_status"),x.get("activity_status"),
        "; ".join(flatten(x.get("buying_signals")))[:200],"; ".join(flatten(x.get("objections")))[:200],
        x.get("authority_concern"),x.get("next_action") or "",(x.get("priority") or "").upper(),link]
    ws2.append(row)
    rr=ws2.max_row
    for c in ws2[rr]:
        c.font=Font(name=FONT,size=9); c.border=border; c.alignment=Alignment(wrap_text=True,vertical="top")
    ws2.cell(rr,7).number_format='#,##0'; ws2.cell(rr,8).number_format='#,##0'
    sc=ws2.cell(rr,4);
    if x["_stage"][1] in STAGE_FILL: sc.fill=PatternFill("solid",fgColor=STAGE_FILL[x["_stage"][1]])
    ac=ws2.cell(rr,17)
    if x.get("ae_summary_status") in AE_FILL: ac.fill=PatternFill("solid",fgColor=AE_FILL[x["ae_summary_status"]])
    vc=ws2.cell(rr,18)
    if x.get("activity_status") in ACT_FILL:
        vc.fill=PatternFill("solid",fgColor=ACT_FILL[x["activity_status"]])
        if x.get("activity_status")=="Dead": vc.font=Font(name=FONT,size=9,color="FFFFFF")
    hc=ws2.cell(rr,24); hc.hyperlink=link; hc.value="Open"; hc.font=Font(name=FONT,size=9,color="0563C1",underline="single")
ws2.freeze_panes="A2"
setw(ws2,{"A":12,"B":30,"C":16,"D":13,"E":9,"F":5,"G":11,"H":12,"I":11,"J":12,"K":7,"L":12,"M":13,"N":11,"O":11,"P":40,"Q":9,"R":9,"S":28,"T":24,"U":6,"V":42,"W":8,"X":8})
for col in ["P","S","T","V"]: pass

# Sheet 3: Missing AE Summary
ws3=wb.create_sheet("Missing AE Summary")
ws3.append(["Deal","AE","Stage","₹","AE Status","Activity","Next Action","HubSpot"]); style_hdr(ws3)
miss=[x for x in deals if x.get("ae_summary_status") in ("No","Partial")]
miss.sort(key=lambda x:({"No":0,"Partial":1}[x["ae_summary_status"]],-x["_inr"]))
for x in miss:
    link=f'https://app.hubspot.com/contacts/1234567/record/0-3/{x["deal_id"]}'
    ws3.append([x.get("dealname"),AE_NAMES.get(x["owner"]),x["_stage"][1],round(x["_inr"]),x.get("ae_summary_status"),x.get("activity_status"),(x.get("next_action") or "")[:160],"Open"])
    rr=ws3.max_row
    for c in ws3[rr]: c.font=Font(name=FONT,size=9); c.border=border; c.alignment=Alignment(wrap_text=True,vertical="top")
    ws3.cell(rr,4).number_format='#,##0'
    if x["ae_summary_status"] in AE_FILL: ws3.cell(rr,5).fill=PatternFill("solid",fgColor=AE_FILL[x["ae_summary_status"]])
    hc=ws3.cell(rr,8); hc.hyperlink=link; hc.font=Font(name=FONT,size=9,color="0563C1",underline="single")
ws3.freeze_panes="A2"; setw(ws3,{"A":30,"B":16,"C":13,"D":12,"E":9,"F":9,"G":50,"H":8})

# Sheet 4: Silent / Dead
ws4=wb.create_sheet("Silent or Dead")
ws4.append(["Deal","AE","Stage","₹","Activity","Days Idle","Last Email","Last Call","Next Action","HubSpot"]); style_hdr(ws4)
sd=[x for x in deals if x.get("activity_status") in ("Silent","Dead")]
sd.sort(key=lambda x:(-x["_stage"][2],-x["_inr"]))
for x in sd:
    link=f'https://app.hubspot.com/contacts/1234567/record/0-3/{x["deal_id"]}'
    ws4.append([x.get("dealname"),AE_NAMES.get(x["owner"]),x["_stage"][1],round(x["_inr"]),x.get("activity_status"),x["_idle"],(x.get("last_email_date") or "")[:10],(x.get("last_call_date") or "")[:10],(x.get("next_action") or "")[:160],"Open"])
    rr=ws4.max_row
    for c in ws4[rr]: c.font=Font(name=FONT,size=9); c.border=border; c.alignment=Alignment(wrap_text=True,vertical="top")
    ws4.cell(rr,4).number_format='#,##0'
    vc=ws4.cell(rr,5)
    if x["activity_status"] in ACT_FILL:
        vc.fill=PatternFill("solid",fgColor=ACT_FILL[x["activity_status"]])
        if x["activity_status"]=="Dead": vc.font=Font(name=FONT,size=9,color="FFFFFF")
    hc=ws4.cell(rr,10); hc.hyperlink=link; hc.font=Font(name=FONT,size=9,color="0563C1",underline="single")
ws4.freeze_panes="A2"; setw(ws4,{"A":30,"B":16,"C":13,"D":12,"E":9,"F":9,"G":11,"H":11,"I":48,"J":8})

# Sheet 5: Stage View (AE x stage matrix)
ws5=wb.create_sheet("Stage View")
stages=["Prospect","Need Analysis","Qualified","Opportunity","Proposal","Negotiation"]
ws5.append(["AE"]+stages+["Total"]); style_hdr(ws5)
for ae in aes:
    ds=[x for x in deals if AE_NAMES.get(x["owner"])==ae]
    row=[ae]+[sum(1 for x in ds if x["_stage"][1]==s) for s in stages]+[len(ds)]
    ws5.append(row)
    for c in ws5[ws5.max_row]: c.font=Font(name=FONT,size=10); c.border=border
ws5.append(["TOTAL"]+[sum(1 for x in deals if x["_stage"][1]==s) for s in stages]+[len(deals)])
for c in ws5[ws5.max_row]: c.font=Font(name=FONT,size=10,bold=True); c.border=border
setw(ws5,{"A":18,"B":11,"C":13,"D":10,"E":11,"F":10,"G":12,"H":8})

# Sheet 6: Notes & Tasks Preview
ws6=wb.create_sheet("Notes & Tasks Preview")
ws6.append(["Deal","AE","Priority","NOTE (will post on deal)","AE TASK (due "+str(AE_DUE)+")","MANAGER TASK (due "+str(MGR_DUE)+")"]); style_hdr(ws6)
def note_plain(d):
    aename=AE_NAMES.get(d["owner"])
    summ=d.get("ae_summary") or ""
    if d.get("ae_summary_status")=="No" or not summ.strip(): summ="AE SUMMARY MISSING - rep must add."
    mgr=(d.get("abhimanyu_note") or "").strip(); nxt=(d.get("next_action") or "").strip()
    s=f"@{aename}\n\nDEAL SUMMARY\n{summ}\n\nABHIMANYU NOTES\n"
    if mgr: s+=mgr+"\n"
    if nxt: s+="Next action: "+nxt+"\n"
    s+=f"\ncc: @{MANAGER_NAME}"
    return s
for x in sorted(deals,key=lambda d:(AE_NAMES.get(d["owner"]),-d["_stage"][2],-d["_inr"])):
    aename=AE_NAMES.get(x["owner"]); pr=(x.get("priority") or "MEDIUM").upper()
    action=(x.get("next_action") or "Review and update deal status.").strip()
    ae_task=f"Action required: {x.get('dealname')}\n\n{action}\n[priority {pr}]"
    mgr_task=f"Review: {x.get('dealname')} — verify AE action\n\nAE ({aename}) was tasked to: {action} Confirm completed. If not, escalate or reassign."
    ws6.append([x.get("dealname"),aename,pr,note_plain(x),ae_task,mgr_task])
    rr=ws6.max_row
    for c in ws6[rr]: c.font=Font(name=FONT,size=9); c.border=border; c.alignment=Alignment(wrap_text=True,vertical="top")
    pc=ws6.cell(rr,3)
    pc.fill=PatternFill("solid",fgColor={"HIGH":"FFC7CE","MEDIUM":"FFEB9C","LOW":"D9D9D9"}.get(pr,"FFFFFF"))
ws6.freeze_panes="A2"; setw(ws6,{"A":28,"B":15,"C":8,"D":60,"E":50,"F":55})

wb.save(OUT_XLSX)
print("workbook saved")
print("payload deals:",len(payload))
from collections import Counter
print("priority:",dict(Counter(p["priority"] for p in payload)))
print("AE summary:",dict(Counter(d.get("ae_summary_status") for d in deals)))
print("activity:",dict(Counter(d.get("activity_status") for d in deals)))
print("total pipeline INR:",round(sum(d["_inr"] for d in deals)))
