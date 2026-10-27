#!/usr/bin/env python3
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

rows=json.load(open('smb_deal_rows.json'))
order={'AE_13':0,'AE_27':1,'SDR_Q':2}
rows.sort(key=lambda r:(order.get(r['owner'],9), -r['stage_order'], -r['amount_inr']))

NAVY='1F3864'; HDR='305496'; WHITE='FFFFFF'
GREEN='C6EFCE'; AMBER='FFEB9C'; LRED='FFC7CE'; RED='FF9999'
GREY='D9D9D9'; LBLUE='DDEBF7'; ZEBRA='F2F7FC'
thin=Side(style='thin',color='BFBFBF')
border=Border(left=thin,right=thin,top=thin,bottom=thin)
F='Arial'
def font(**k): return Font(name=F,**k)

wb=Workbook()

# ---------- Deals sheet ----------
ws=wb.active; ws.title='Deals'
cols=['Owner','Deal','Stage','Amount','Cur','Amount (INR)','Created','Close','Last Activity','Days Idle',
      'Emails In','Emails Out','Calls Conn','Calls Att','Last Email','Last Call','Notes','AE Summary Status','Activity','AE Deal Summary','Recommended Next Action','HubSpot']
ws.append(cols)
for c in range(1,len(cols)+1):
    cell=ws.cell(1,c); cell.font=font(bold=True,color=WHITE,size=10)
    cell.fill=PatternFill('solid',fgColor=HDR); cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); cell.border=border
ws.freeze_panes='A2'
act_fill={'Active':GREEN,'One-way':AMBER,'Silent':LRED,'Dead':RED}
ae_fill={'Yes':GREEN,'Partial':AMBER,'No':RED}
stage_fill={'5-Negotiation':GREEN,'4-Proposal':AMBER,'3-Opportunity':LRED,'2-Need Analysis':LBLUE,'1-Prospect':GREY}
r=2
for d in rows:
    ws.append([d['owner'],d['dealname'],d['stage'],d['amount'],d['currency'],d['amount_inr'],d['createdate'],d['closedate'],
        d['last_activity'],d['days_idle'],d['emails_in'],d['emails_out'],d['calls_connected'],d['calls_attempted'],
        d['last_email'],d['last_call'],d['notes'],d['ae_status'],d['activity'],d['ae_summary'],d['next_action'],'Open'])
    ws.cell(r,4).number_format='#,##0'
    ws.cell(r,6).number_format='#,##0'
    ws.cell(r,22).hyperlink=d['url']; ws.cell(r,22).font=font(color='0563C1',underline='single',size=9)
    ws.cell(r,3).fill=PatternFill('solid',fgColor=stage_fill.get(d['stage'],WHITE))
    ws.cell(r,18).fill=PatternFill('solid',fgColor=ae_fill.get(d['ae_status'],WHITE))
    ws.cell(r,19).fill=PatternFill('solid',fgColor=act_fill.get(d['activity'],WHITE))
    for c in range(1,len(cols)+1):
        cell=ws.cell(r,c); cell.border=border
        if cell.font is None or cell.font.name!=F:
            cell.font=font(size=9)
        cell.alignment=Alignment(vertical='top',wrap_text=(c in (2,20,21)),horizontal='left' if c in (1,2,3,20,21) else 'center')
        if c in (20,21): cell.font=font(size=9)
        if r%2==0 and c not in (3,18,19):
            if cell.fill.fgColor.rgb in (None,'90000001'):
                cell.fill=PatternFill('solid',fgColor=ZEBRA)
    r+=1
widths=[16,26,15,10,5,12,11,11,12,8,8,8,8,8,11,11,6,11,10,60,46,7]
for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
for rr in range(2,r): ws.row_dimensions[rr].height=70

# ---------- Summary sheet ----------
s=wb.create_sheet('Summary',0)
s['A1']='SMB Hunting Team — Deal Summary'; s['A1'].font=font(bold=True,size=16,color=NAVY)
s['A2']='Owners: AE_13, AE_27, SDR_Q  ·  Open deals, stages 1–5  ·  Source: engagement cache (last sync 2026-07-01)'
s['A2'].font=font(size=9,italic=True,color='595959')
def hdrrow(ws,row,labels,start=1,fill=HDR):
    for i,l in enumerate(labels):
        c=ws.cell(row,start+i); c.value=l; c.font=font(bold=True,color=WHITE,size=10)
        c.fill=PatternFill('solid',fgColor=fill); c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); c.border=border

# by-owner scorecard
s['A4']='Pipeline & Coverage by Owner'; s['A4'].font=font(bold=True,size=12,color=NAVY)
hdrrow(s,5,['Owner','Deals','Pipeline (INR)','Active','One-way','Silent','Dead','AE Summary: Yes','Partial','No'])
owners=['AE_13','AE_27','SDR_Q']
rr=6
for o in owners:
    od=[d for d in rows if d['owner']==o]
    def cnt(key,val): return sum(1 for d in od if d[key]==val)
    s.append([o,len(od),sum(d['amount_inr'] for d in od),cnt('activity','Active'),cnt('activity','One-way'),
              cnt('activity','Silent'),cnt('activity','Dead'),cnt('ae_status','Yes'),cnt('ae_status','Partial'),cnt('ae_status','No')])
    rr+=1
# totals row with formulas
s.cell(rr,1,'TOTAL').font=font(bold=True)
for col in range(2,11):
    L=get_column_letter(col); s.cell(rr,col,f'=SUM({L}6:{L}{rr-1})').font=font(bold=True)
for row in range(6,rr+1):
    s.cell(row,3).number_format='#,##0'
    for col in range(1,11):
        c=s.cell(row,col); c.border=border
        if c.font.name!=F: c.font=font(size=10)
        c.alignment=Alignment(horizontal='left' if col==1 else 'center')
    if row==rr:
        for col in range(1,11): s.cell(row,col).fill=PatternFill('solid',fgColor=LBLUE)

# critical attention list
crit=[d for d in rows if (d['activity'] in ('Silent','Dead') and d['stage_order']>=4) or d['ae_status']=='No']
crit.sort(key=lambda d:-d['amount_inr'])
base=rr+3
s.cell(base-1,1,'Critical Attention — advanced-stage deals gone quiet or missing context').font=font(bold=True,size=12,color='C00000')
hdrrow(s,base,['Owner','Deal','Stage','Amount (INR)','Activity','Days Idle','Why flagged'],fill='C00000')
r2=base+1
for d in crit:
    why=[]
    if d['activity'] in ('Silent','Dead'): why.append(f"{d['activity'].lower()} at {d['stage']}")
    if d['ae_status']=='No': why.append('no summary logged')
    s.cell(r2,1,d['owner']); s.cell(r2,2,d['dealname']); s.cell(r2,3,d['stage'])
    s.cell(r2,4,d['amount_inr']).number_format='#,##0'; s.cell(r2,5,d['activity'])
    s.cell(r2,6,d['days_idle']); s.cell(r2,7,'; '.join(why))
    s.cell(r2,5).fill=PatternFill('solid',fgColor=act_fill.get(d['activity'],WHITE))
    for col in range(1,8):
        c=s.cell(r2,col); c.border=border
        if c.font.name!=F: c.font=font(size=9)
        c.alignment=Alignment(horizontal='left' if col in(1,2,3,7) else 'center',wrap_text=col==7)
    r2+=1
for i,w in enumerate([16,30,15,14,10,10,34],1): s.column_dimensions[get_column_letter(i)].width=w
s.freeze_panes='A6'

wb.save('SMB_Hunting_Deal_Summary.xlsx')
print('saved rows',len(rows),'critical',len(crit))
